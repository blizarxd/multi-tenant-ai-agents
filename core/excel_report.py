"""
Generador de reportes Excel para historial médico de pacientes.
Requiere: pip install openpyxl
"""

import io
import json
import logging
from datetime import datetime

logger = logging.getLogger(__name__)

# Colores de triage
_TRIAGE_COLORS = {
    "URGENTE":   "FFCCCC",  # rojo claro
    "PRIORIDAD": "FFF0B0",  # amarillo claro
    "RUTINA":    "CCFFCC",  # verde claro
}
_HEADER_COLOR   = "1E5AA0"  # azul oscuro
_SUBHEADER_COLOR = "D0E4FF"  # azul claro
_ALT_ROW_COLOR  = "F5F8FF"  # azul muy claro (filas pares)


def generate_patient_excel(tenant_id: str, clinica: str, doctor: str, memory) -> bytes:
    """
    Genera un Excel con dos hojas:
      1. Resumen de pacientes
      2. Historial detallado de citas
    Retorna bytes listos para enviar por Telegram o subir a Drive.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("Instala openpyxl: pip install openpyxl")

    wb = Workbook()

    def _header_font(white=True):
        return Font(bold=True, color="FFFFFF" if white else "1E5AA0", size=10)

    def _fill(hex_color: str):
        return PatternFill("solid", fgColor=hex_color)

    def _border():
        thin = Side(style="thin", color="CCCCCC")
        return Border(left=thin, right=thin, top=thin, bottom=thin)

    def _center():
        return Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _left():
        return Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ═══ HOJA 1: Resumen de pacientes ════════════════════════════
    ws1 = wb.active
    ws1.title = "Pacientes"
    ws1.row_dimensions[1].height = 30

    # Título
    ws1.merge_cells("A1:H1")
    ws1["A1"] = f"{clinica} — {doctor} | Historial de Pacientes | {datetime.now().strftime('%d/%m/%Y')}"
    ws1["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    ws1["A1"].fill = _fill(_HEADER_COLOR)
    ws1["A1"].alignment = _center()

    # Cabecera de tabla
    headers = ["#", "Nombre", "Teléfono", "Primera Visita", "Última Visita",
               "Total Eventos", "Última Cita", "Último Motivo"]
    ws1.row_dimensions[2].height = 20
    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=2, column=col, value=h)
        cell.font = _header_font()
        cell.fill = _fill(_HEADER_COLOR)
        cell.alignment = _center()
        cell.border = _border()

    patients = memory.get_all_patients()
    for i, pat in enumerate(patients, 1):
        row = i + 2
        fill = _fill(_ALT_ROW_COLOR) if i % 2 == 0 else PatternFill()

        # Obtener última cita
        citas = memory.get_events_for_patient(pat["telegram_id"], "cita")
        last_cita_fecha = ""
        last_cita_motivo = ""
        if citas:
            p = citas[0]["payload"]
            last_cita_fecha  = p.get("fecha", citas[0].get("timestamp", "")[:10])
            last_cita_motivo = p.get("motivo", "")

        vals = [
            i,
            pat.get("nombre") or "Sin nombre",
            pat.get("telefono") or "—",
            (pat.get("first_seen") or "")[:10],
            (pat.get("last_seen")  or "")[:10],
            pat.get("total_eventos", 0),
            last_cita_fecha,
            last_cita_motivo,
        ]
        for col, val in enumerate(vals, 1):
            cell = ws1.cell(row=row, column=col, value=val)
            cell.fill = fill
            cell.border = _border()
            cell.alignment = _center() if col in (1, 4, 5, 6) else _left()

    # Anchos de columna
    col_widths = [5, 28, 18, 14, 14, 10, 22, 40]
    for col, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(col)].width = w

    ws1.freeze_panes = "A3"

    # ═══ HOJA 2: Historial detallado de citas ════════════════════
    ws2 = wb.create_sheet("Historial Citas")
    ws2.row_dimensions[1].height = 30

    ws2.merge_cells("A1:J1")
    ws2["A1"] = f"{clinica} — Historial Clínico Detallado | {datetime.now().strftime('%d/%m/%Y')}"
    ws2["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    ws2["A1"].fill = _fill(_HEADER_COLOR)
    ws2["A1"].alignment = _center()

    headers2 = ["Fecha Cita", "Paciente", "Teléfono", "Motivo", "Triage",
                "Servicio", "Precio USD", "Notas", "Tipo Evento", "Fecha Registro"]
    ws2.row_dimensions[2].height = 20
    for col, h in enumerate(headers2, 1):
        cell = ws2.cell(row=2, column=col, value=h)
        cell.font = _header_font()
        cell.fill = _fill(_HEADER_COLOR)
        cell.alignment = _center()
        cell.border = _border()

    row2 = 3
    for pat in patients:
        events = memory.get_events_for_patient(pat["telegram_id"])
        cita_events = [e for e in events if e["event_type"] in ("cita", "urgencia", "cita_cancelada")]
        for ev in reversed(cita_events):  # cronológico
            p = ev["payload"]
            triage = p.get("triage", "RUTINA").upper()
            triage_color = _TRIAGE_COLORS.get(triage, "FFFFFF")

            vals = [
                p.get("fecha", (ev.get("timestamp") or "")[:10]),
                p.get("paciente") or pat.get("nombre") or "—",
                p.get("telefono") or pat.get("telefono") or "—",
                p.get("motivo") or p.get("sintoma") or "—",
                triage,
                p.get("servicio", "—"),
                p.get("precio_usd") or p.get("precio", "—"),
                p.get("notas", "—"),
                ev["event_type"],
                (ev.get("timestamp") or "")[:16],
            ]
            for col, val in enumerate(vals, 1):
                cell = ws2.cell(row=row2, column=col, value=val)
                cell.border = _border()
                cell.alignment = _center() if col in (5, 6, 7, 9) else _left()
                if col == 5:  # columna Triage — color por nivel
                    cell.fill = _fill(triage_color)
                    cell.font = Font(bold=True, size=9)
                elif row2 % 2 == 0:
                    cell.fill = _fill(_ALT_ROW_COLOR)

            ws2.row_dimensions[row2].height = 16
            row2 += 1

    col_widths2 = [20, 28, 18, 40, 12, 22, 10, 40, 16, 16]
    for col, w in enumerate(col_widths2, 1):
        ws2.column_dimensions[get_column_letter(col)].width = w

    ws2.freeze_panes = "A3"

    # ── Guardar en bytes ──────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
